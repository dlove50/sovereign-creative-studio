from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path("outputs/canadian-personal-data-removal")
OUT.mkdir(parents=True, exist_ok=True)
DOCX = OUT / "Canadian_Personal_Data_Removal_Guide.docx"

BLUE = "1F4E78"
PALE = "EAF2F8"
GOLD = "FFF2CC"
RED = "FCE4D6"
GRAY = "5B6573"

def font(run, size=11, bold=False, color="222222", name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc.get_or_add_tcPr()
    tcMar = tc.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tc.append(tcMar)
    for tag, value in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")
        tcMar.append(node)

def set_table_widths(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in("w:tblW")
    if tblW is None:
        tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths))); tblW.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(width)); grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tcW = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tcW is None:
                tcW = OxmlElement("w:tcW"); cell._tc.get_or_add_tcPr().append(tcW)
            tcW.set(qn("w:w"), str(width)); tcW.set(qn("w:type"), "dxa")
            margins(cell)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)
    return p

def callout(doc, label, text, fill=PALE):
    t = doc.add_table(rows=1, cols=1)
    set_table_widths(t, [9360])
    shade(t.cell(0,0), fill)
    p = t.cell(0,0).paragraphs[0]
    r = p.add_run(label + " "); font(r, 10.5, True, BLUE)
    r = p.add_run(text); font(r, 10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.49)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"; normal.font.size = Pt(11)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.18
for name, size, before, after in (("Heading 1",16,18,8),("Heading 2",13,14,6),("Heading 3",11.5,10,4)):
    s = styles[name]; s.font.name="Aptos Display"; s.font.size=Pt(size); s.font.bold=True
    s.font.color.rgb=RGBColor.from_string(BLUE)
    s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
    s.paragraph_format.keep_with_next=True

header = sec.header.paragraphs[0]
header.text = "PRIVATE WORKING GUIDE  •  ONTARIO, CANADA"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(header.runs[0], 8.5, True, GRAY)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(footer.add_run("Keep the guide and tracker on a trusted local device. Review law and site procedures before acting."), 8, False, GRAY)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
r = p.add_run("Canadian Personal-Data\nRemoval Guide"); font(r, 25, True, BLUE, "Aptos Display")
p = doc.add_paragraph()
r = p.add_run("A cautious discovery, request, and recheck workflow for an Ontario resident"); font(r, 12, False, GRAY)
p.paragraph_format.space_after = Pt(14)
callout(doc, "Purpose.", "Find exposed personal information, document it privately, and prepare lawful site-specific requests. This guide does not promise complete erasure and is not legal advice.")

add_heading(doc, "The safe operating rule", 1)
doc.add_paragraph("Start with discovery only. Do not create accounts, submit forms, bypass CAPTCHAs, upload identification, open verification links, or press final submission buttons during discovery. Each of those actions requires your review.")
add_bullet(doc, "Use the least information needed to find a listing.")
add_bullet(doc, "Prefer the site’s official privacy or opt-out process.")
add_bullet(doc, "Do not claim a law applies until the organization, activity, and jurisdiction have been checked.")
add_bullet(doc, "Save evidence and confirmations, but avoid duplicating sensitive data unnecessarily.")

add_heading(doc, "1. Prepare a minimal private search profile", 1)
doc.add_paragraph("Do not put a street address, date of birth, SIN, government ID, financial data, passwords, or relatives in the initial profile. Use placeholders until you deliberately fill the file on your own device.")
code = """{
  "full_name": "YOUR_FULL_NAME",
  "aliases": [],
  "city": "YOUR_CITY",
  "province": "Ontario",
  "country": "Canada",
  "search_email": "YOUR_EMAIL",
  "search_phone": "YOUR_PHONE"
}"""
t=doc.add_table(rows=1,cols=1); set_table_widths(t,[9360]); shade(t.cell(0,0),"F3F5F7")
p=t.cell(0,0).paragraphs[0]; r=p.add_run(code); font(r,9.5,False,"202124","Consolas")

add_heading(doc, "Working Windows setup", 2)
doc.add_paragraph("Run these native PowerShell commands from a folder you control. They create a private working directory and placeholder profile; they do not install software or send data anywhere.")
ps = """$privacyRoot = Join-Path $env:USERPROFILE 'Documents\\Privacy-Removal'
New-Item -ItemType Directory -Path $privacyRoot -Force | Out-Null
$profilePath = Join-Path $privacyRoot 'search-profile.json'
@'
{
  "full_name": "YOUR_FULL_NAME",
  "aliases": [],
  "city": "YOUR_CITY",
  "province": "Ontario",
  "country": "Canada",
  "search_email": "YOUR_EMAIL",
  "search_phone": "YOUR_PHONE"
}
'@ | Set-Content -LiteralPath $profilePath -Encoding utf8
Get-Item -LiteralPath $profilePath | Select-Object FullName, Length, LastWriteTime"""
t=doc.add_table(rows=1,cols=1); set_table_widths(t,[9360]); shade(t.cell(0,0),"F3F5F7")
p=t.cell(0,0).paragraphs[0]; r=p.add_run(ps); font(r,8.5,False,"202124","Consolas")
callout(doc, "Avoid unsupported tooling.", "Do not rely on a reversible redaction/reconstruction workflow unless the tool’s current documentation explicitly supports it. Redaction is not a backup strategy.", GOLD)

add_heading(doc, "2. Discover and classify listings", 1)
doc.add_paragraph("Search a few identifiers at a time. Record only what is needed to relocate and assess each result. Store screenshots locally when they are genuinely useful evidence.")
for item in [
    "Exact name in quotation marks plus city or province.",
    "Known email address or phone number, one at a time.",
    "People-search sites, business directories, public social profiles, and search-engine results.",
    "Official privacy policy, privacy contact, and opt-out page for each organization."
]: add_bullet(doc,item)

add_heading(doc, "Risk scale", 2)
tab=doc.add_table(rows=4,cols=3); set_table_widths(tab,[1300,3100,4960])
rows=[("Risk","Examples","Priority"),("High","Home address, phone, birth date, relatives, or combined identifiers","Review promptly; preserve evidence; use official removal route"),("Medium","Personal email, employment history, detailed profile","Confirm accuracy and exposure; prepare a tailored request"),("Low","Ordinary professional or public reference","Document; remove only if inaccurate, unwanted, or unnecessary")]
for i,row in enumerate(rows):
    for j,val in enumerate(row):
        tab.cell(i,j).text=val
        if i==0: shade(tab.cell(i,j),BLUE)
        for run in tab.cell(i,j).paragraphs[0].runs: font(run,9.2,i==0,"FFFFFF" if i==0 else "222222")

add_heading(doc, "3. Choose the right request", 1)
doc.add_paragraph("For many Ontario private-sector commercial activities, PIPEDA is the starting point. Ontario health records may instead fall under PHIPA; provincial and municipal institutions use Ontario public-sector access/privacy laws. Foreign organizations may also provide contractual opt-outs or rights under another jurisdiction.")
callout(doc, "Important legal limit.", "PIPEDA provides access and correction rights and permits withdrawal of consent subject to legal or contractual restrictions and reasonable notice. Deletion can follow in some circumstances, but there is no universal promise that every record must be erased.", GOLD)

add_heading(doc, "Template A — simple opt-out or suppression request", 2)
doc.add_paragraph("""Subject: Request to remove or suppress my public listing

Hello Privacy Team,

I am the individual shown at [LISTING URL]. Please remove or suppress this public listing and stop displaying the associated personal information in public search results.

Information sufficient to locate the record:
Name: [NAME]
City/province: [CITY], Ontario
Listing URL: [URL]

Please tell me if you need additional information, why it is necessary, how it will be protected, and how long it will be retained. Please confirm when the listing is removed or suppressed and identify any steps I must complete myself.

Thank you,
[NAME]""")

add_heading(doc, "Template B — access and correction request", 2)
doc.add_paragraph("""Subject: Request for access to and correction of personal information

Hello Privacy Officer,

I am requesting access to the personal information your organization holds about me. Please tell me:
1. what personal information you hold;
2. how it has been used;
3. the source of the information; and
4. the organizations to which it has been disclosed, where applicable.

The following information appears inaccurate or incomplete: [DESCRIBE].
The correct information is: [CORRECTION].

Please explain your identity-verification requirements before I send additional documents. I prefer a method that collects the least additional personal information.

Please acknowledge receipt and advise of the expected response date.

Thank you,
[NAME]""")

add_heading(doc, "Template C — withdrawal of consent and deletion inquiry", 2)
doc.add_paragraph("""Subject: Withdrawal of consent and request to assess deletion

Hello Privacy Officer,

To the extent your organization relies on my consent for the continued collection, use, or disclosure of the personal information associated with [ACCOUNT/LISTING], I withdraw that consent, subject to any applicable legal or contractual restrictions and reasonable notice.

Please explain:
1. the consequences of withdrawal;
2. any information you must retain and the legal or contractual basis;
3. whether the remaining information can be deleted, de-identified, or suppressed; and
4. when the action will be completed.

Please do not request more identifying information than is necessary to verify and process this request.

Thank you,
[NAME]""")

add_heading(doc, "4. Manual review and submission gate", 1)
for item in [
    "Confirm the listing is yours and the URL is correct.",
    "Read the current opt-out/privacy instructions and note the organization’s jurisdiction.",
    "Remove unnecessary identifiers from the request.",
    "Complete CAPTCHA, login, ID upload, email link, and final submission yourself.",
    "Record the submission date, method, confirmation number, and expected follow-up date.",
    "Keep any identity image out of the tracker; record only that verification occurred."
]: add_bullet(doc,item)
callout(doc, "Identity-document warning.", "If a site asks for ID, first ask whether a masked copy or another verification method is accepted. Cover fields the site does not need. Do not upload a SIN card.", RED)

add_heading(doc, "5. Follow up and escalate carefully", 1)
doc.add_paragraph("Organizations subject to PIPEDA generally must respond to an access request within 30 calendar days, with limited extensions and notice requirements. A site-specific opt-out may use a different service target. Track the actual promise made.")
for item in [
    "Send a polite follow-up after the stated deadline.",
    "Ask for written reasons if access, correction, suppression, or deletion is refused.",
    "Use the organization’s internal privacy complaint process first where practical.",
    "If the issue is within PIPEDA and remains unresolved, review the Office of the Privacy Commissioner of Canada complaint process. Different regulators may apply to health, government, or foreign organizations."
]: add_bullet(doc,item)

add_heading(doc, "6. Recheck without overclaiming success", 1)
doc.add_paragraph("Recheck high-risk listings after 7–14 days, again after 30 days, and then quarterly or on a schedule appropriate to the exposure. Weekly checks can be noisy and may unnecessarily repeat sensitive searches.")
for item in [
    "Open the direct listing URL and repeat the minimum search that originally found it.",
    "Check whether search-engine snippets or cached results lag behind the source site.",
    "Record reappearance as a new event linked to the original request.",
    "Document lawful public records or other unremovable material instead of repeatedly submitting the same request.",
    "Success means prioritized listings were removed or suppressed, errors corrected, confirmations saved, and reappearance checks completed—not that the entire Internet was certified clean."
]: add_bullet(doc,item)

add_heading(doc, "Discovery-only AI prompt", 1)
prompt = """TASK: Help me identify publicly accessible websites that may contain my personal information so I can prepare lawful privacy or opt-out requests.

JURISDICTION: Ontario, Canada.

SEARCH INFORMATION: Use only the name, aliases, city/province, email, and phone that I explicitly provide.

SAFETY: Do not expose complete identifiers in the report. Do not use a SIN, birth date, government ID, financial data, passwords, or relatives. Do not create accounts, submit forms, bypass CAPTCHAs, upload ID, open verification links, or click final submission. Do not claim a law applies without verifying jurisdiction and requirements. Flag all actions requiring my approval.

OUTPUT COLUMNS: Website; Listing URL; Information Exposed; Risk; Official Opt-Out URL; Jurisdiction; Recommended Action; Manual Verification Likely; Status; Date Checked.

Stop after the discovery report. Do not contact any organization."""
t=doc.add_table(rows=1,cols=1); set_table_widths(t,[9360]); shade(t.cell(0,0),"F3F5F7")
p=t.cell(0,0).paragraphs[0]; r=p.add_run(prompt); font(r,9,False,"202124","Consolas")

add_heading(doc, "Source notes (checked 28 July 2026)", 1)
sources = [
("Office of the Privacy Commissioner of Canada — Accessing personal information held by businesses","https://www.priv.gc.ca/en/privacy-topics/accessing-personal-information/api_bus/"),
("OPC — Responding to access requests under PIPEDA","https://www.priv.gc.ca/en/privacy-topics/accessing-personal-information/obligations-for-organizations/02_05_d_54_ati_02/"),
("OPC — PIPEDA Principle 3: Consent","https://www.priv.gc.ca/en/privacy-topics/privacy-laws-in-canada/the-personal-information-protection-and-electronic-documents-act-pipeda/p_principle/principles/p_consent/"),
("OPC — Case Summary 2017-005: deletion after withdrawal of consent","https://www.priv.gc.ca/en/opc-actions-and-decisions/investigations/investigations-into-businesses/2017/pipeda-2017-005/"),
("Ontario IPC — Access and privacy in private-sector businesses","https://www.ipc.on.ca/en/media-centre/blog/access-and-privacy-private-sector-businesses"),
("Ontario IPC — What we do and do not oversee","https://www.ipc.on.ca/en/about-us/role-and-mandate/what-we-do-vs-what-we-dont-do"),
]
for title,url in sources:
    p=doc.add_paragraph(style="List Bullet")
    p.add_run(title + ": ")
    p.add_run(url)

doc.core_properties.title = "Canadian Personal-Data Removal Guide"
doc.core_properties.subject = "Cautious personal-data discovery and removal workflow for an Ontario resident"
doc.core_properties.author = "Codex"
doc.save(DOCX)
print(DOCX.resolve())
