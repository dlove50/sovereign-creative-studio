from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Derived from this file's location, matching check_transcription_deps.py and
# local_launcher.py. A hardcoded absolute path broke this script on Linux and on
# any Windows machine where the repo lives somewhere else (e.g. the laptop kit).
ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "outputs" / "how-to-guides" / "Baidu_Unlimited_OCR_Installation_and_Sovereign_Integration_Guide.docx"
ICON = ROOT / "assets" / "ocr-launcher" / "unlimited-ocr-icon.png"

NAVY = "12162B"; VIOLET = "7C5CFC"; CYAN = "2DD4FF"; INK = "20243A"
MUTED = "667085"; PALE = "F2F1FF"; CODE_BG = "171A2B"; WHITE = "FFFFFF"
GREEN = "067647"; GREEN_BG = "ECFDF3"; GOLD = "8A6116"; GOLD_BG = "FFFAEB"
RED = "B42318"; RED_BG = "FEF3F2"; BORDER = "D9DCE7"

def set_cell_fill(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn("w:shd"))
    if shd is None: shd = OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"), color)

def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None: tcMar = OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None: node = OxmlElement(f"w:{m}"); tcMar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")

def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            tcPr = row.cells[idx]._tc.get_or_add_tcPr(); tcW = tcPr.find(qn("w:tcW"))
            if tcW is None: tcW = OxmlElement("w:tcW"); tcPr.append(tcW)
            tcW.set(qn("w:w"), str(int(width * 1440))); tcW.set(qn("w:type"), "dxa")
            set_cell_margins(row.cells[idx])
    tblPr = table._tbl.tblPr; tblW = tblPr.find(qn("w:tblW"))
    if tblW is None: tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(sum(widths)*1440))); tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None: tblInd = OxmlElement("w:tblInd"); tblPr.append(tblInd)
    tblInd.set(qn("w:w"), "120"); tblInd.set(qn("w:type"), "dxa")

def shade_paragraph(p, color, border=None):
    pPr = p._p.get_or_add_pPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), color); pPr.append(shd)
    if border:
        pBdr = OxmlElement("w:pBdr"); left = OxmlElement("w:left")
        left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "18"); left.set(qn("w:space"), "8"); left.set(qn("w:color"), border)
        pBdr.append(left); pPr.append(pBdr)

def set_font(run, name="Aptos", size=10.5, color=INK, bold=False, italic=False):
    run.font.name = name; run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name); run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string(color); run.bold = bold; run.italic = italic

doc = Document()
sec = doc.sections[0]; sec.top_margin = Inches(.78); sec.bottom_margin = Inches(.75); sec.left_margin = Inches(.88); sec.right_margin = Inches(.88)
sec.header_distance = Inches(.35); sec.footer_distance = Inches(.35)

styles = doc.styles
normal = styles["Normal"]; normal.font.name = "Aptos"; normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.18
for name, size, color, before, after in (("Title", 30, NAVY, 0, 8), ("Subtitle", 14, MUTED, 0, 16), ("Heading 1", 18, NAVY, 18, 9), ("Heading 2", 14, VIOLET, 14, 6), ("Heading 3", 11.5, NAVY, 10, 4)):
    st = styles[name]; st.font.name = "Aptos Display" if name != "Normal" else "Aptos"; st.font.size = Pt(size); st.font.color.rgb = RGBColor.from_string(color); st.font.bold = name != "Subtitle"
    st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after); st.paragraph_format.keep_with_next = True

header = sec.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run("UNLIMITED OCR  /  IMPLEMENTATION MANUAL"); set_font(r, size=8, color=MUTED, bold=True)
footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("Sovereign Creative Studio  •  Local-first AI workflow  •  2026"); set_font(r, size=8, color=MUTED)

def p(text="", style=None, bold_prefix=None):
    para = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        a = para.add_run(bold_prefix); set_font(a, bold=True)
        b = para.add_run(text[len(bold_prefix):]); set_font(b)
    else: para.add_run(text)
    return para

def bullet(text):
    para = doc.add_paragraph(style="List Bullet"); para.paragraph_format.left_indent = Inches(.25); para.paragraph_format.first_line_indent = Inches(-.17); para.paragraph_format.space_after = Pt(3); para.add_run(text); return para

def step(num, title, body):
    para = doc.add_paragraph(); para.paragraph_format.space_before = Pt(8); para.paragraph_format.space_after = Pt(3)
    badge = para.add_run(f"{num:02d}  "); set_font(badge, size=11, color=VIOLET, bold=True)
    rr = para.add_run(title); set_font(rr, size=11, color=NAVY, bold=True)
    p(body)

def callout(label, text, kind="info"):
    colors = {"info": (PALE, VIOLET), "success": (GREEN_BG, GREEN), "warn": (GOLD_BG, GOLD), "risk": (RED_BG, RED)}
    bg, accent = colors[kind]
    para = doc.add_paragraph(); para.paragraph_format.left_indent = Inches(.1); para.paragraph_format.right_indent = Inches(.1); para.paragraph_format.space_before = Pt(6); para.paragraph_format.space_after = Pt(8)
    shade_paragraph(para, bg, accent); a = para.add_run(f"  {label.upper()}  "); set_font(a, size=9.5, color=accent, bold=True); b = para.add_run(text + "  "); set_font(b, size=9.5, color=INK)

def code(title, text, language=""):
    para = doc.add_paragraph(); para.paragraph_format.space_before = Pt(5); para.paragraph_format.space_after = Pt(8); para.paragraph_format.left_indent = Inches(.08); para.paragraph_format.right_indent = Inches(.08)
    shade_paragraph(para, CODE_BG)
    h = para.add_run(f"  {title}{'  ·  '+language if language else ''}\n"); set_font(h, name="Cascadia Mono", size=8.5, color=CYAN, bold=True)
    rr = para.add_run(text.rstrip() + "  "); set_font(rr, name="Cascadia Mono", size=8, color=WHITE)

def page_break(): doc.add_page_break()

# Cover
doc.add_paragraph("IMPLEMENTATION MANUAL", style=None).runs[0].font.color.rgb = RGBColor.from_string(VIOLET)
if ICON.exists():
    pp = doc.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER; pp.paragraph_format.space_before = Pt(20); pp.paragraph_format.space_after = Pt(18); pp.add_run().add_picture(str(ICON), width=Inches(2.05))
title = doc.add_paragraph(style="Title"); title.alignment = WD_ALIGN_PARAGRAPH.CENTER; title.add_run("Baidu Unlimited-OCR")
sub = doc.add_paragraph(style="Subtitle"); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER; sub.add_run("From original Hugging Face model to a polished local OCR Studio")
rule = doc.add_paragraph(); rule.alignment = WD_ALIGN_PARAGRAPH.CENTER; rr = rule.add_run("━━━━━━━━━━━━━━━━━━━━"); set_font(rr, size=10, color=CYAN)
meta = doc.add_paragraph(); meta.alignment = WD_ALIGN_PARAGRAPH.CENTER; meta.paragraph_format.space_before = Pt(14)
for text, bold in (("DUAL HOW-TO GUIDE\n", True), ("PART I  •  Install the original full BF16 model\n", False), ("PART II  •  Integrate it into Sovereign Creative Studio\n\n", False), ("Windows 11  |  NVIDIA RTX GPU  |  Docker Desktop  |  vLLM  |  Next.js", False)):
    rr = meta.add_run(text); set_font(rr, size=10.5 if not bold else 11, color=NAVY if bold else MUTED, bold=bold)
callout("Outcome", "A private, local OCR system that accepts images and PDFs, preserves document structure as Markdown, and launches from a custom desktop shortcut.", "success")
page_break()

# Front matter
p("How to use this manual", "Heading 1")
p("This document is both a clean-room installation guide and an implementation record. Part I can be followed on a new Windows PC. Part II assumes the model API from Part I is working and shows how to add the OCR experience to an existing Next.js application such as Sovereign Creative Studio.")
p("Scope and tested configuration", "Heading 2")
t = doc.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = "Table Grid"; set_table_widths(t, [1.65, 4.85])
for c, text in zip(t.rows[0].cells, ["Component", "Reference configuration"]): set_cell_fill(c, NAVY); c.paragraphs[0].add_run(text); set_font(c.paragraphs[0].runs[0], size=9, color=WHITE, bold=True)
for a,b in [("Host", "Windows 11, PowerShell 5.1+"),("GPU", "NVIDIA GeForce RTX 5070 Ti, 16 GB VRAM"),("Runtime", "Docker Desktop with WSL 2 GPU passthrough"),("Model", "baidu/Unlimited-OCR, original BF16 Safetensors, approximately 3B parameters"),("Serving", "Dedicated vLLM Unlimited-OCR image, OpenAI-compatible API"),("Application", "Next.js 14 / React 18 / TypeScript / Tailwind CSS")]:
    cells=t.add_row().cells; cells[0].text=a; cells[1].text=b; cells[0].vertical_alignment=cells[1].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
callout("Terminology", "“Original model” means the files served directly from the official baidu/Unlimited-OCR repository without GGUF conversion or weight quantization. The vLLM runtime is separate from the model weights.")
p("Contents", "Heading 2")
for text in ["Part I — Install and run Baidu Unlimited-OCR", "1. Architecture and prerequisites", "2. Docker and GPU verification", "3. Official runtime and original model deployment", "4. API verification and OCR tests", "5. Operations, security, and troubleshooting", "Part II — Integrate Unlimited-OCR into Sovereign Creative Studio", "6. Product architecture", "7. Frontend dependencies and navigation", "8. Image/PDF ingestion and OCR requests", "9. Results, history, and exports", "10. Background launcher and custom desktop shortcut", "11. Build, validation, and acceptance checklist", "Appendices — Complete reference scripts and command index"]: bullet(text)
page_break()

# Part I
p("PART I", "Heading 1"); p("Install and run the original Baidu Unlimited-OCR model", "Title")
p("This part reproduces the full installation path used for the working local system. It deliberately avoids Ollama conversions and downloads the public Baidu checkpoint into a persistent Hugging Face cache.")
callout("Why not Ollama?", "Ollama requires architecture support and usually a GGUF or repackaged model. That is useful for convenience, but it is not the untouched original Hugging Face BF16 checkpoint. vLLM runs the official model format and the model-specific decoding recipe.", "warn")

p("1. Architecture and prerequisites", "Heading 1")
p("The local request path", "Heading 2")
code("Runtime flow", "Browser or API client\n  → http://127.0.0.1:8000/v1/chat/completions\n  → Docker Desktop / WSL 2\n  → vLLM dedicated Unlimited-OCR runtime\n  → baidu/Unlimited-OCR BF16 weights\n  → NVIDIA GPU\n  → Markdown + grounding tokens")
p("Minimum practical prerequisites", "Heading 2")
for item in ["Windows 11 with virtualization enabled.", "Docker Desktop configured for the WSL 2 backend.", "A supported NVIDIA GPU; the official vLLM recipe states 8 GB VRAM is sufficient for BF16 inference.", "Current NVIDIA Windows driver with WSL/Docker GPU support.", "At least 25–35 GB free disk space for the CUDA/vLLM image, model cache, and working headroom.", "Internet access for the initial Docker and Hugging Face downloads."]: bullet(item)
code("Inventory checks", "docker --version\nwsl --version\nnvidia-smi", "PowerShell")

p("2. Docker and GPU verification", "Heading 1")
step(1, "Install or update WSL", "Run the following from an elevated PowerShell window, then restart Windows if requested.")
code("WSL setup", "wsl --install -d Ubuntu\nwsl --update", "PowerShell (Administrator)")
step(2, "Install Docker Desktop", "Install Docker Desktop for Windows. In Settings, enable the WSL 2 based engine and WSL integration for Ubuntu.")
step(3, "Start Docker Desktop", "Wait until Docker Desktop reports that its engine is running.")
code("Docker engine check", "docker info --format 'Server={{.ServerVersion}} OS={{.OperatingSystem}}'", "PowerShell")
step(4, "Verify NVIDIA GPU passthrough", "This test downloads a small CUDA base image and executes nvidia-smi inside Linux. The GPU name and VRAM should match the Windows host.")
code("Container GPU test", "docker run --rm --gpus all nvidia/cuda:13.0.0-base-ubuntu24.04 nvidia-smi", "PowerShell")
callout("Pass condition", "The output lists the NVIDIA GPU inside the container. Do not continue if Docker reports that no GPU driver or compatible device is available.", "success")

p("3. Install the official runtime and original model", "Heading 1")
step(1, "Pull the dedicated vLLM image", "Use the model-specific tag published for Unlimited-OCR. This image carries the new architecture support and required logits processor.")
code("Runtime download", "docker pull vllm/vllm-openai:unlimited-ocr", "PowerShell")
step(2, "Create a persistent Hugging Face cache", "A named volume keeps the original checkpoint after container recreation or upgrades.")
code("Persistent cache", "docker volume create baidu-unlimited-ocr-cache", "PowerShell")
step(3, "Launch the service", "Bind the API to 127.0.0.1 so it is not exposed to the local network. The model identifier explicitly points to Baidu’s official repository.")
code("Production-style local launch", "docker run -d --name baidu-unlimited-ocr --restart unless-stopped --gpus all --ipc=host `\n  -p 127.0.0.1:8000:8000 `\n  -v baidu-unlimited-ocr-cache:/root/.cache/huggingface `\n  vllm/vllm-openai:unlimited-ocr `\n  baidu/Unlimited-OCR `\n  --trust-remote-code `\n  --logits_processors vllm.model_executor.models.unlimited_ocr:NGramPerReqLogitsProcessor `\n  --no-enable-prefix-caching `\n  --mm-processor-cache-gb 0 `\n  --gpu-memory-utilization 0.85", "PowerShell")
callout("Important", "The first start downloads approximately 6.2 GiB of original checkpoint data and compiles GPU kernels. It can take several minutes. Later starts reuse the cache.", "info")
p("What each non-default option does", "Heading 2")
t=doc.add_table(rows=1, cols=2); t.style="Table Grid"; set_table_widths(t,[2.25,4.25])
for c,x in zip(t.rows[0].cells,["Option","Purpose"]): set_cell_fill(c,NAVY); c.paragraphs[0].add_run(x); set_font(c.paragraphs[0].runs[0],size=9,color=WHITE,bold=True)
for a,b in [("--trust-remote-code","Loads Baidu’s custom Transformers architecture code."),("--logits_processors …","Registers the model-specific no-repeat n-gram processor."),("--no-enable-prefix-caching","Matches the OCR workload and official recipe."),("--mm-processor-cache-gb 0","Disables unnecessary multimodal processor caching."),("--gpu-memory-utilization 0.85","Leaves host/display headroom on a 16 GB desktop GPU."),("--restart unless-stopped","Restarts automatically with Docker unless intentionally stopped.")]: cells=t.add_row().cells; cells[0].text=a; cells[1].text=b

p("4. Verify the API and perform OCR", "Heading 1")
step(1, "Watch startup logs", "Wait for “Application startup complete” and the server route list.")
code("Live logs", "docker logs -f baidu-unlimited-ocr\n# Press Ctrl+C to leave the log viewer; the container keeps running.", "PowerShell")
step(2, "Confirm the exact model", "A successful response should contain id and root values equal to baidu/Unlimited-OCR and max_model_len equal to 32768.")
code("Model readiness", "Invoke-RestMethod http://127.0.0.1:8000/v1/models", "PowerShell")
step(3, "OCR an image with PowerShell", "The request embeds the image as a data URL. The literal <image> prefix, skip_special_tokens=false, and n-gram arguments are required for reliable output.")
code("End-to-end PowerShell client", "$imagePath = 'C:\\Documents\\sample.png'\n$base64 = [Convert]::ToBase64String([IO.File]::ReadAllBytes($imagePath))\n$payload = @{\n  model = 'baidu/Unlimited-OCR'\n  messages = @(@{ role = 'user'; content = @(\n    @{ type = 'text'; text = '<image>document parsing.' },\n    @{ type = 'image_url'; image_url = @{ url = \"data:image/png;base64,$base64\" } }\n  )})\n  max_tokens = 8192\n  temperature = 0.0\n  skip_special_tokens = $false\n  vllm_xargs = @{ ngram_size = 35; window_size = 128 }\n} | ConvertTo-Json -Depth 10 -Compress\n\n$result = Invoke-RestMethod `\n  -Uri 'http://127.0.0.1:8000/v1/chat/completions' `\n  -Method Post -ContentType 'application/json' `\n  -Body $payload -TimeoutSec 600\n$result.choices[0].message.content", "PowerShell")
step(4, "OCR an image with Python", "This client uses the OpenAI-compatible endpoint but keeps all inference local.")
code("Python API client", "import base64\nfrom pathlib import Path\nfrom openai import OpenAI\n\nimage_path = Path(r\"C:\\Documents\\sample.png\")\nmime = \"image/png\"\ndata = base64.b64encode(image_path.read_bytes()).decode(\"ascii\")\n\nclient = OpenAI(api_key=\"EMPTY\", base_url=\"http://127.0.0.1:8000/v1\", timeout=600)\nresponse = client.chat.completions.create(\n    model=\"baidu/Unlimited-OCR\",\n    messages=[{\"role\": \"user\", \"content\": [\n        {\"type\": \"text\", \"text\": \"<image>document parsing.\"},\n        {\"type\": \"image_url\", \"image_url\": {\"url\": f\"data:{mime};base64,{data}\"}},\n    ]}],\n    max_tokens=8192, temperature=0.0,\n    extra_body={\n        \"skip_special_tokens\": False,\n        \"vllm_xargs\": {\"ngram_size\": 35, \"window_size\": 128},\n    },\n)\nprint(response.choices[0].message.content)", "Python")

p("5. Operations, security, and troubleshooting", "Heading 1")
code("Daily operations", "docker start baidu-unlimited-ocr\ndocker stop baidu-unlimited-ocr\ndocker ps --filter name=baidu-unlimited-ocr\ndocker logs --tail 100 baidu-unlimited-ocr", "PowerShell")
p("Common problems", "Heading 2")
for title, body in [("The URL is not recognized in PowerShell", "A URL typed by itself is treated as a command. Use Start-Process 'http://127.0.0.1:8000/docs'."),("The output is empty", "Ensure the prompt begins with the literal <image> token and skip_special_tokens is false."),("The output repeats coordinate tokens", "Ensure the server was started with NGramPerReqLogitsProcessor and each request supplies ngram_size=35 and window_size=128."),("Out of GPU memory", "Close other GPU-heavy applications, reduce --gpu-memory-utilization cautiously, or shorten max_tokens/context usage."),("Docker cannot see the GPU", "Update the NVIDIA driver, confirm WSL 2 mode, restart Docker Desktop, and rerun the CUDA container test."),("Strange banner characters", "This is PowerShell/terminal encoding of vLLM’s block-art banner; it does not affect inference.")]:
    pp=doc.add_paragraph(); a=pp.add_run(title+": "); set_font(a,bold=True,color=NAVY); b=pp.add_run(body); set_font(b)
page_break()

# Part II
p("PART II", "Heading 1"); p("Integrate Unlimited-OCR into Sovereign Creative Studio", "Title")
p("This part turns the developer API into a polished local product: drag-and-drop input, PDF rendering, multi-page progress, Markdown/plain/raw views, local history, downloads, health state, and one-click startup.")

p("6. Product architecture", "Heading 1")
code("Application flow", "Desktop shortcut\n  → hidden PowerShell readiness launcher\n  → Docker Desktop + baidu-unlimited-ocr\n  → Next.js frontend at /ocr-studio\n  → browser converts PDF pages to JPEG data URLs\n  → sequential POST /v1/chat/completions\n  → clean Markdown + raw grounding-token view\n  → localStorage history + copy/download")
p("Design decisions", "Heading 2")
for item in ["Keep the official vLLM backend independent from the app frontend.", "Send files directly from the browser to 127.0.0.1:8000; no cloud relay or project database is required.", "Render PDFs in the browser, avoiding a second server-side conversion service.", "Process pages sequentially for predictable GPU memory use and clear progress feedback.", "Retain raw model output for diagnostics while presenting cleaned Markdown by default.", "Store only compact completed-result history in browser localStorage; source documents are not persisted."]: bullet(item)

p("7. Frontend dependencies and navigation", "Heading 1")
step(1, "Install document-view dependencies", "pdfjs-dist renders PDF pages in the browser. react-markdown and remark-gfm display headings, tables, lists, and code blocks.")
code("Frontend dependencies", "cd frontend\nnpm install pdfjs-dist@4.10.38 react-markdown@9.0.1 remark-gfm@4.0.0\nCopy-Item node_modules\\pdfjs-dist\\build\\pdf.worker.min.mjs public\\pdf.worker.min.mjs", "PowerShell")
step(2, "Add the route", "Create a small route component that hosts the client-side OCR workspace.")
code("frontend/src/app/ocr-studio/page.tsx", "import { OcrStudioView } from \"@/components/ocr-studio-view\";\n\nexport default function OcrStudioPage() {\n  return <OcrStudioView />;\n}", "TypeScript / React")
step(3, "Add navigation", "Import ScanText from lucide-react and add the route to the creator navigation array.")
code("Navigation item", "{ href: \"/ocr-studio\", label: \"OCR Studio\", icon: ScanText },", "TypeScript")

p("8. Image/PDF ingestion and OCR requests", "Heading 1")
p("Browser-side PDF rendering", "Heading 2")
code("PDF pages to data URLs", "const pdfjs = await import(\"pdfjs-dist\");\npdfjs.GlobalWorkerOptions.workerSrc = \"/pdf.worker.min.mjs\";\nconst bytes = new Uint8Array(await file.arrayBuffer());\nconst pdf = await pdfjs.getDocument({ data: bytes }).promise;\n\nfor (let pageNumber = 1; pageNumber <= pdf.numPages; pageNumber++) {\n  const page = await pdf.getPage(pageNumber);\n  const viewport = page.getViewport({ scale: 2 });\n  const canvas = document.createElement(\"canvas\");\n  canvas.width = viewport.width; canvas.height = viewport.height;\n  const context = canvas.getContext(\"2d\");\n  await page.render({ canvasContext: context!, viewport }).promise;\n  const dataUrl = canvas.toDataURL(\"image/jpeg\", 0.94);\n  // Queue dataUrl as one OCR page.\n}", "TypeScript")
p("Model health check", "Heading 2")
code("Local health state", "const response = await fetch(\n  \"http://127.0.0.1:8000/v1/models\",\n  { signal: AbortSignal.timeout(5000) },\n);\nconst online = response.ok;", "TypeScript")
p("Sequential page recognition", "Heading 2")
code("OCR request", "const response = await fetch(\n  \"http://127.0.0.1:8000/v1/chat/completions\",\n  {\n    method: \"POST\",\n    headers: { \"Content-Type\": \"application/json\" },\n    body: JSON.stringify({\n      model: \"baidu/Unlimited-OCR\",\n      temperature: 0,\n      max_tokens: 8192,\n      messages: [{ role: \"user\", content: [\n        { type: \"text\", text: \"<image>document parsing.\" },\n        { type: \"image_url\", image_url: { url: page.dataUrl } },\n      ]}],\n      skip_special_tokens: false,\n      vllm_xargs: { ngram_size: 35, window_size: 128 },\n    }),\n  },\n);\nif (!response.ok) throw new Error(await response.text());\nconst payload = await response.json();\nconst raw = payload.choices?.[0]?.message?.content ?? \"\";", "TypeScript")
callout("CORS", "vLLM permits the local frontend origin by default in this configuration. Verify with an OPTIONS request if the browser reports a cross-origin error.")

p("9. Clean results, history, and exports", "Heading 1")
p("Remove grounding markers", "Heading 2")
code("Reference cleaner", "function cleanOutput(raw: string) {\n  return raw\n    .replace(/<\\|ref\\|>/g, \"\")\n    .replace(/<\\|\\/ref\\|>/g, \"\")\n    .replace(/<\\|det\\|>.*?<\\|\\/det\\|>/g, \"\")\n    .trim();\n}", "TypeScript")
callout("Production note", "A line-aware cleaner is preferable because <|det|> markers can prefix useful text on the same line. Preserve a raw-output view for debugging and future parser improvements.", "warn")
p("Local history", "Heading 2")
code("Browser-only persistence", "const HISTORY_KEY = \"sovereign-ocr-history-v1\";\nconst next = [completedItem, ...history].slice(0, 12);\nlocalStorage.setItem(HISTORY_KEY, JSON.stringify(next));", "TypeScript")
p("Download helper", "Heading 2")
code("Markdown/text download", "function downloadText(content: string, filename: string) {\n  const url = URL.createObjectURL(new Blob([content], { type: \"text/markdown\" }));\n  const anchor = document.createElement(\"a\");\n  anchor.href = url; anchor.download = filename; anchor.click();\n  URL.revokeObjectURL(url);\n}", "TypeScript")

p("10. One-click background launcher and desktop shortcut", "Heading 1")
p("The launcher is responsible for orchestration, not inference. It starts Docker Desktop when necessary, waits for the Docker engine, starts the named model container, starts the production frontend only when absent, polls both URLs, and opens the OCR page.")
code("PowerShell launcher — essential logic", "$appRoot = 'C:\\Path\\To\\sovereign-creative-studio'\n$dockerDesktop = 'C:\\Program Files\\Docker\\Docker\\Docker Desktop.exe'\n$ocrHealth = 'http://127.0.0.1:8000/v1/models'\n$studioUrl = 'http://127.0.0.1:3000/ocr-studio'\n\ndocker info *> $null\nif ($LASTEXITCODE -ne 0) {\n  Start-Process -FilePath $dockerDesktop -WindowStyle Hidden\n  for ($i=0; $i -lt 60; $i++) {\n    Start-Sleep -Seconds 2; docker info *> $null\n    if ($LASTEXITCODE -eq 0) { break }\n  }\n}\n\ndocker start baidu-unlimited-ocr *> $null\nif (-not (Test-Url $studioUrl)) {\n  Start-Process powershell.exe `\n    -ArgumentList '-NoProfile','-ExecutionPolicy','Bypass','-File',\n      (Join-Path $appRoot 'start-frontend.ps1') `\n    -WorkingDirectory $appRoot -WindowStyle Hidden\n}\n\nfor ($i=0; $i -lt 90; $i++) {\n  if ((Test-Url $ocrHealth) -and (Test-Url $studioUrl)) { break }\n  Start-Sleep -Seconds 2\n}\nStart-Process $studioUrl", "PowerShell")
code("BAT wrapper", "@echo off\nsetlocal EnableExtensions\nset \"LAUNCHER=C:\\Path\\To\\tools\\Start-Unlimited-OCR.ps1\"\nstart \"Unlimited OCR\" /min powershell.exe -NoProfile `\n  -ExecutionPolicy Bypass -WindowStyle Hidden -File \"%LAUNCHER%\"\nexit /b 0", "BAT")
p("Create a custom Windows shortcut", "Heading 2")
code("Shortcut creation", "$desktop = [Environment]::GetFolderPath('Desktop')\n$shell = New-Object -ComObject WScript.Shell\n$shortcut = $shell.CreateShortcut((Join-Path $desktop 'Unlimited OCR.lnk'))\n$shortcut.TargetPath = (Join-Path $desktop 'Launch Unlimited OCR.bat')\n$shortcut.WorkingDirectory = 'C:\\Path\\To\\sovereign-creative-studio'\n$shortcut.IconLocation = 'C:\\Path\\To\\unlimited-ocr.ico,0'\n$shortcut.Description = 'Start Baidu Unlimited OCR and open OCR Studio'\n$shortcut.WindowStyle = 7\n$shortcut.Save()", "PowerShell")
p("Generate a multi-resolution ICO with Python", "Heading 2")
code("PNG to Windows icon", "from PIL import Image\nfrom pathlib import Path\n\nsrc = Path(r\"C:\\Path\\To\\unlimited-ocr-icon.png\")\nout = Path(r\"C:\\Path\\To\\unlimited-ocr.ico\")\nimage = Image.open(src).convert(\"RGBA\")\nimage.save(out, format=\"ICO\", sizes=[\n    (16,16), (24,24), (32,32), (48,48),\n    (64,64), (128,128), (256,256),\n])", "Python")

p("11. Build, validation, and acceptance", "Heading 1")
code("Frontend verification", "cd frontend\nnpm run typecheck\nnpm run build", "PowerShell")
p("Acceptance checklist", "Heading 2")
for item in ["Docker Desktop can start from a stopped state.", "The container reports baidu/Unlimited-OCR at /v1/models.", "The API is bound only to 127.0.0.1:8000.", "The OCR Studio route loads at 127.0.0.1:3000/ocr-studio.", "Drag-and-drop works for PNG, JPG, WEBP, and PDF.", "A multi-page PDF processes sequentially and reports progress.", "Markdown tables and headings render correctly.", "Plain and raw output views work.", "Copy and download produce the expected text.", "Recent result history persists after browser refresh.", "The desktop shortcut starts services in the background and opens the correct route.", "Stopping the OCR container releases GPU memory."]: bullet("☐ " + item)
callout("Security acceptance", "Keep the API on localhost, do not add a public tunnel, and do not expose port 8000 to the LAN without authentication and transport security.", "risk")

p("Appendix A — Exact installed project locations", "Heading 1")
for a,b in [("Model service","Docker container: baidu-unlimited-ocr"),("Model cache","Docker volume: baidu-unlimited-ocr-cache"),("OCR page","frontend/src/app/ocr-studio/page.tsx"),("OCR interface","frontend/src/components/ocr-studio-view.tsx"),("Navigation","frontend/src/lib/navigation.ts"),("PDF worker","frontend/public/pdf.worker.min.mjs"),("PowerShell launcher","tools/Start-Unlimited-OCR.ps1"),("BAT launcher","Launch-Unlimited-OCR.bat"),("Icon assets","assets/ocr-launcher/"),("Desktop shortcut","Unlimited OCR.lnk")]:
    pp=doc.add_paragraph(); x=pp.add_run(a+": "); set_font(x,bold=True,color=NAVY); y=pp.add_run(b); set_font(y,name="Cascadia Mono",size=9)

p("Appendix B — Recovery procedure", "Heading 1")
step(1,"Restart Docker Desktop","Exit Docker Desktop fully, open it again, and wait for the engine.")
step(2,"Restart the OCR container","Run docker stop followed by docker start for baidu-unlimited-ocr.")
step(3,"Verify model identity","Call /v1/models and confirm baidu/Unlimited-OCR.")
step(4,"Rebuild the frontend","Run npm run typecheck and npm run build in the frontend directory.")
step(5,"Start the frontend","Run the project’s start-frontend.ps1 production launcher.")
step(6,"Retest one image","Use the PowerShell or Python client from Part I before troubleshooting the UI.")

p("Appendix C — Source references", "Heading 1")
for text in ["Official model: https://huggingface.co/baidu/Unlimited-OCR", "Official source repository: https://github.com/baidu/Unlimited-OCR", "Official vLLM recipe: https://recipes.vllm.ai/baidu/Unlimited-OCR", "vLLM runtime image: vllm/vllm-openai:unlimited-ocr"]: bullet(text)
p("End of manual")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.core_properties.title = "Baidu Unlimited-OCR Installation and Sovereign Integration Guide"
doc.core_properties.subject = "Dual how-to manual for original model deployment and local application integration"
doc.core_properties.author = "Sovereign Creative Studio"
doc.save(OUT)
print(OUT)
